"""Lecteur de fichiers Excel (XLS, XLSX, XLSM).

Utilise openpyxl (XLSX) et xlrd (XLS) pour lire les classeurs Excel
et les convertir en lignes list[list[str]] compatibles avec le pipeline CSV.

Dépendances :
- openpyxl : pip install openpyxl  (XLSX, XLSM)
- xlrd     : pip install xlrd      (XLS ancien format)
"""

import logging
from datetime import datetime, date, time
from pathlib import Path
from typing import Generator, Optional

logger = logging.getLogger(__name__)


def _cell_to_str(value) -> str:
    """Convertit une valeur de cellule Excel en string."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime('%d/%m/%Y')
        return value.strftime('%d/%m/%Y %H:%M:%S')
    if isinstance(value, date):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, time):
        return value.strftime('%H:%M:%S')
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value)
    return str(value).strip()


def read_xlsx(
    filepath: Path,
    sheet_name: Optional[str] = None,
    sheet_index: int = 0,
) -> tuple[list[str], list[list[str]], int]:
    """Lit un fichier XLSX/XLSM et retourne header + toutes les lignes.

    Utilise le mode read_only pour le streaming mémoire.

    Args:
        filepath: Chemin du fichier.
        sheet_name: Nom de la feuille (prioritaire sur sheet_index).
        sheet_index: Index de la feuille (0-based).

    Returns:
        Tuple (header, rows, total_lines).
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl est requis pour lire les fichiers XLSX. "
            "Installez-le avec : pip install openpyxl"
        )

    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise ValueError(
                f"Feuille '{sheet_name}' introuvable. "
                f"Feuilles disponibles : {wb.sheetnames}"
            )
        ws = wb[sheet_name]
    else:
        ws = wb.worksheets[min(sheet_index, len(wb.worksheets) - 1)]

    logger.info(f"Lecture XLSX : {filepath.name} → feuille '{ws.title}'")

    header = []
    rows = []
    for i, row in enumerate(ws.iter_rows()):
        cells = [_cell_to_str(cell.value) for cell in row]
        if i == 0:
            header = cells
        else:
            rows.append(cells)

    wb.close()

    total_lines = len(rows) + 1  # +1 pour le header
    logger.info(f"XLSX lu : {total_lines} lignes, {len(header)} colonnes")
    return header, rows, total_lines


def read_xlsx_chunks(
    filepath: Path,
    chunk_size: int = 10_000,
    sheet_name: Optional[str] = None,
    sheet_index: int = 0,
    skip_header: bool = False,
) -> Generator[list[list[str]], None, None]:
    """Lit un fichier XLSX par chunks en streaming.

    Args:
        filepath: Chemin du fichier.
        chunk_size: Nombre de lignes par chunk.
        sheet_name: Nom de la feuille.
        sheet_index: Index de la feuille.
        skip_header: Si True, ne pas inclure le header.

    Yields:
        Chunks de lignes (chaque ligne = list[str]).
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl requis : pip install openpyxl")

    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.worksheets[min(sheet_index, len(wb.worksheets) - 1)]

    chunk = []
    for i, row in enumerate(ws.iter_rows()):
        if i == 0 and skip_header:
            continue
        cells = [_cell_to_str(cell.value) for cell in row]
        chunk.append(cells)
        if len(chunk) >= chunk_size:
            yield chunk
            chunk = []

    if chunk:
        yield chunk
    wb.close()


def read_xls(
    filepath: Path,
    sheet_name: Optional[str] = None,
    sheet_index: int = 0,
) -> tuple[list[str], list[list[str]], int]:
    """Lit un fichier XLS (ancien format) et retourne header + lignes.

    Args:
        filepath: Chemin du fichier.
        sheet_name: Nom de la feuille.
        sheet_index: Index de la feuille.

    Returns:
        Tuple (header, rows, total_lines).
    """
    try:
        import xlrd
    except ImportError:
        raise ImportError(
            "xlrd est requis pour lire les fichiers XLS. "
            "Installez-le avec : pip install xlrd"
        )

    wb = xlrd.open_workbook(str(filepath))

    if sheet_name:
        try:
            ws = wb.sheet_by_name(sheet_name)
        except xlrd.XLRDError:
            raise ValueError(
                f"Feuille '{sheet_name}' introuvable. "
                f"Feuilles disponibles : {wb.sheet_names()}"
            )
    else:
        ws = wb.sheet_by_index(min(sheet_index, wb.nsheets - 1))

    logger.info(f"Lecture XLS : {filepath.name} → feuille '{ws.name}'")

    header = []
    rows = []

    for row_idx in range(ws.nrows):
        cells = []
        for col_idx in range(ws.ncols):
            cell = ws.cell(row_idx, col_idx)
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    dt_tuple = xlrd.xldate_as_tuple(cell.value, wb.datemode)
                    if dt_tuple[3:] == (0, 0, 0):
                        cells.append(f"{dt_tuple[2]:02d}/{dt_tuple[1]:02d}/{dt_tuple[0]:04d}")
                    else:
                        cells.append(
                            f"{dt_tuple[2]:02d}/{dt_tuple[1]:02d}/{dt_tuple[0]:04d} "
                            f"{dt_tuple[3]:02d}:{dt_tuple[4]:02d}:{dt_tuple[5]:02d}"
                        )
                except Exception:
                    cells.append(str(cell.value))
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                cells.append('TRUE' if cell.value else 'FALSE')
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                if cell.value == int(cell.value):
                    cells.append(str(int(cell.value)))
                else:
                    cells.append(str(cell.value))
            elif cell.ctype == xlrd.XL_CELL_EMPTY:
                cells.append('')
            else:
                cells.append(str(cell.value).strip())

        if row_idx == 0:
            header = cells
        else:
            rows.append(cells)

    total_lines = len(rows) + 1
    logger.info(f"XLS lu : {total_lines} lignes, {len(header)} colonnes")
    return header, rows, total_lines


def list_sheets(filepath: Path) -> list[str]:
    """Liste les noms de feuilles d'un fichier Excel.

    Args:
        filepath: Chemin du fichier XLS ou XLSX.

    Returns:
        Liste des noms de feuilles.
    """
    ext = filepath.suffix.lower()
    if ext in ('.xlsx', '.xlsm'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(filepath), read_only=True)
            names = wb.sheetnames
            wb.close()
            return names
        except ImportError:
            raise ImportError("openpyxl requis : pip install openpyxl")
    elif ext == '.xls':
        try:
            import xlrd
            wb = xlrd.open_workbook(str(filepath))
            return wb.sheet_names()
        except ImportError:
            raise ImportError("xlrd requis : pip install xlrd")
    return []


def read_excel(
    filepath: Path,
    sheet_name: Optional[str] = None,
    sheet_index: int = 0,
) -> tuple[list[str], list[list[str]], int]:
    """Lit un fichier Excel (XLS ou XLSX) automatiquement.

    Args:
        filepath: Chemin du fichier.
        sheet_name: Nom de la feuille (optionnel).
        sheet_index: Index de la feuille (défaut: 0).

    Returns:
        Tuple (header, rows, total_lines).
    """
    ext = filepath.suffix.lower()
    if ext in ('.xlsx', '.xlsm'):
        return read_xlsx(filepath, sheet_name, sheet_index)
    elif ext == '.xls':
        return read_xls(filepath, sheet_name, sheet_index)
    else:
        raise ValueError(f"Extension non supportée : {ext}. Utilisez .xls, .xlsx ou .xlsm")
